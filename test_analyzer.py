from io import BytesIO

import pymupdf
from docx import Document

from analyzer import extract_resume_text

# Create and test a sample PDF
pdf = pymupdf.open()
page = pdf.new_page()
page.insert_text(
    (72, 72),
    "Sample Resume\nPython, Machine Learning, Flask"
)

pdf_bytes = pdf.tobytes()
pdf.close()

pdf_text = extract_resume_text(
    pdf_bytes,
    "sample_resume.pdf"
)

print("PDF extraction successful:")
print(pdf_text)

# Create and test a sample DOCX
doc = Document()
doc.add_heading("Sample Resume", level=1)
doc.add_paragraph("Python, Machine Learning, Flask")

buffer = BytesIO()
doc.save(buffer)

docx_text = extract_resume_text(
    buffer.getvalue(),
    "sample_resume.docx"
)

print("\nDOCX extraction successful:")
print(docx_text)

print("\nBoth tests completed successfully!")