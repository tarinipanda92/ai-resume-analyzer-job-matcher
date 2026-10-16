
from pathlib import Path
from io import BytesIO

import pymupdf
from docx import Document


def extract_resume_text(file_bytes: bytes, filename: str) -> str:
    """Extract text from an uploaded PDF or DOCX resume."""

    if not file_bytes:
        raise ValueError("The uploaded file is empty.")

    extension = Path(filename).suffix.lower()

    if extension == ".pdf":
        with pymupdf.open(
            stream=file_bytes,
            filetype="pdf"
        ) as pdf:
            text = "\n".join(
                page.get_text("text")
                for page in pdf
            )

    elif extension == ".docx":
        document = Document(BytesIO(file_bytes))
        parts = []

        # Extract text from paragraphs.
        for paragraph in document.paragraphs:
            if paragraph.text.strip():
                parts.append(paragraph.text.strip())

        # Extract text from tables.
        for table in document.tables:
            for row in table.rows:
                cells = [
                    cell.text.strip()
                    for cell in row.cells
                    if cell.text.strip()
                ]

                if cells:
                    parts.append(" | ".join(cells))

        text = "\n".join(parts)

    else:
        raise ValueError(
            "Unsupported file type. Please upload a PDF or DOCX file."
        )

    # Remove empty lines and surrounding whitespace.
    text = "\n".join(
        line.strip()
        for line in text.splitlines()
        if line.strip()
    )

    if not text:
        raise ValueError(
            "No readable text found. The file may be empty "
            "or contain scanned images only."
        )

    return text
